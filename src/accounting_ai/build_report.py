"""Build financial statements from classified trial balance items."""

from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from accounting_ai.types.ledger import AccountType, EntrySide, LedgerItem


def build_report(
    items: list[LedgerItem],
    output_file: Path,
) -> None:
    """Create a new workbook with both statements and save it to output_file."""
    grouped = {account_type: [] for account_type in AccountType}
    for item in items:
        if item.account_type is None:
            raise ValueError(f"Account {item.name!r} has no account type.")
        normal_side = (
            EntrySide.debit
            if item.account_type in (AccountType.asset, AccountType.expenses)
            else EntrySide.credit
        )
        amount = Decimal(str(item.value))
        if not amount.is_finite():
            raise ValueError(f"Account {item.name!r} has a non-finite amount.")
        if item.entry_side != normal_side:
            amount = -amount
        grouped[item.account_type].append((item.name, amount))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def statement(title):
        ws = wb.create_sheet(title)
        ws.append([title, "Amount"])
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 20
        return ws

    def total_row(ws, label, amount):
        ws.append([label, amount])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    def section(ws, title, account_type):
        ws.append([title])
        for name, amount in grouped[account_type]:
            ws.append([name, amount])
        total = sum((amount for _, amount in grouped[account_type]), Decimal(0))
        total_row(ws, f"Total {title.lower()}", total)
        ws.append([])
        return total

    pnl = statement("Profit and Loss")
    revenue = section(pnl, "Revenue", AccountType.revenue)
    expenses = section(pnl, "Expenses", AccountType.expenses)
    profit = revenue - expenses
    total_row(pnl, "Profit / (loss) for the period", profit)

    position = statement("Statement of Financial Position")
    assets = section(position, "Assets", AccountType.asset)
    liabilities = section(position, "Liabilities", AccountType.liability)
    capital = section(position, "Capital", AccountType.capital)
    position.append(["Profit / (loss) for the period", profit])
    total_row(position, "Total equity", capital + profit)
    total_row(position, "Total liabilities and equity", liabilities + capital + profit)
    total_row(position, "Balance check (should be zero)", assets - liabilities - capital - profit)

    for ws in (pnl, position):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = '#,##0.00;(#,##0.00);0.00'

    try:
        wb.save(output_file)
    finally:
        wb.close()
